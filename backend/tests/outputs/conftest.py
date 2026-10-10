"""Fixtures for the output tests (Story 7.2): spaces and auth, plus a temporary files
directory, and helpers to build XLSX/DOCX/CSV content and create outputs."""

import io
import uuid
from collections.abc import AsyncIterator, Iterator
from contextlib import asynccontextmanager
from pathlib import Path

import httpx
import pytest
from docx import Document
from openpyxl import Workbook  # type: ignore[import-untyped]

from app.auth.models import User
from app.auth.service import SessionUser
from app.core.db import transaction
from app.outputs import service
from app.outputs.models import OutputKind, OutputOrigin
from app.outputs.settings import get_outputs_settings
from app.spaces.access import AccessContext
from app.spaces.deps import build_access_context
from tests.audit.conftest import audit_key
from tests.auth.conftest import auth_env, auth_settings_env, csrf_headers, login, make_client
from tests.crypto.conftest import kek_bytes
from tests.spaces.conftest import spaces_db

__all__ = ["audit_key", "auth_env", "auth_settings_env", "kek_bytes", "spaces_db"]


@pytest.fixture(autouse=True)
def files_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Iterator[Path]:
    path = tmp_path / "files"
    monkeypatch.setenv("APP_OUTPUTS_FILES_DIR", str(path))
    get_outputs_settings.cache_clear()
    yield path
    get_outputs_settings.cache_clear()


async def context_of(user: User) -> AccessContext:
    return await build_access_context(SessionUser(user.id, user.username, user.role, user.id))


@asynccontextmanager
async def signed_in(username: str) -> AsyncIterator[httpx.AsyncClient]:
    async with make_client() as client:
        assert (await login(client, username)).status_code == 200
        client.headers.update(csrf_headers(client))
        yield client


async def chunks(data: bytes, size: int = 64 * 1024) -> AsyncIterator[bytes]:
    for start in range(0, len(data), size):
        yield data[start : start + size]


def xlsx_bytes(rows: list[list[object]], *, sheets: int = 1) -> bytes:
    workbook = Workbook()
    for index in range(sheets):
        sheet = workbook.active if index == 0 else workbook.create_sheet(f"Sheet{index + 1}")
        for row in rows:
            sheet.append(row)
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def docx_bytes(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


async def create(
    user: User,
    space_id: uuid.UUID,
    data: bytes,
    *,
    kind: OutputKind = OutputKind.CSV,
    title: str = "Margins table",
    version_of_id: uuid.UUID | None = None,
) -> service.OutputView:
    ctx = await context_of(user)
    async with transaction(context=ctx.rls_settings()) as tx:
        return await service.create_output(
            tx,
            ctx,
            space_id=space_id,
            kind=kind,
            origin=OutputOrigin.OPERATIONS,
            content=chunks(data),
            title=title,
            source_document_ids=[uuid.uuid4()],
            version_of_id=version_of_id,
            ops_log=[{"op": "filter"}],
        )


def private_space(ctx: AccessContext) -> uuid.UUID:
    (space_id,) = ctx.space_ids
    return space_id
