"""AC #1 (RLS), AC #2 (encrypted storage), AC #4 (links and downloads), AC #5 (manage)."""

import hashlib
import io
import uuid
from datetime import timedelta
from pathlib import Path

import httpx
import pytest
from docx import Document
from openpyxl import load_workbook  # type: ignore[import-untyped]
from sqlalchemy import text

from app.audit.testing import assert_no_content
from app.audit.types import KeyedHash
from app.auth.models import Role
from app.core.db import transaction
from app.core.timeutil import utcnow
from app.crypto.fields import IntegrityError
from app.crypto.files import CHUNK_SIZE, decrypt_stream
from app.crypto.keys import KeyDestroyed
from app.outputs import service
from app.outputs.models import OutputKind
from tests.auth.conftest import AuthEnv, audit_rows
from tests.conftest import PgDatabase
from tests.outputs.conftest import (
    chunks,
    context_of,
    create,
    docx_bytes,
    private_space,
    signed_in,
    xlsx_bytes,
)
from tests.spaces.conftest import as_migrator, as_user, make_user, make_workspace

pytestmark = [pytest.mark.db, pytest.mark.usefixtures("spaces_db")]


def _files(root: Path) -> dict[Path, bytes]:
    """Every file under ``root`` and its content."""
    return {p: p.read_bytes() for p in root.rglob("*") if p.is_file()}


PLANTED = b"Planted,Revenue\r\nFalcon-7731,1250\r\n"
OUTPUTS = "/api/v1/outputs"


# --- AC #2 --------------------------------------------------------------------------------


async def test_file_is_encrypted_on_disk_and_decrypts(files_dir: Path) -> None:
    user = await make_user("anika")
    ctx = await context_of(user)
    space_id = private_space(ctx)
    created = await create(user, space_id, PLANTED * 1000)

    path = files_dir / str(space_id) / str(created.id)
    stored = path.read_bytes()
    assert b"Falcon-7731" not in stored
    assert [p.name for p in path.parent.iterdir()] == [str(created.id)]  # no temp files
    for content in _files(files_dir).values():
        assert b"Planted" not in content
    ref = service.KeyRef(space_id, service.OBJECT_TYPE, created.id)
    plain = b"".join([c async for c in decrypt_stream(ref, chunks(stored), name="content")])
    assert plain == PLANTED * 1000
    assert created.size_bytes == len(PLANTED) * 1000


async def test_expiry_follows_space_retention() -> None:
    user = await make_user("bo")
    ctx = await context_of(user)
    created = await create(user, private_space(ctx), PLANTED)
    assert timedelta(days=29, hours=23) < created.expires_at - utcnow() <= timedelta(days=30)


async def test_create_needs_access_and_cleans_up_on_failure(files_dir: Path) -> None:
    user = await make_user("cleo")
    other = await make_user("dev")
    other_space = private_space(await context_of(other))
    with pytest.raises(service.SpaceNotAccessibleError):
        await create(user, other_space, PLANTED)
    with pytest.raises(ValueError, match="title"):
        await create(user, private_space(await context_of(user)), PLANTED, title="  ")
    assert _files(files_dir) == {}


# --- AC #1 --------------------------------------------------------------------------------


async def test_owner_only_inside_a_workspace_and_membership_needed(pg: PgDatabase) -> None:
    a = await make_user("alice")
    b = await make_user("bilal")
    lead = await make_user("lead")
    workspace = await make_workspace("Falcon", lead, a, b)
    created = await create(a, workspace, PLANTED)
    await service.issue_download_link(await context_of(a), created.id)

    for table in ("outputs", "download_tokens"):
        assert await as_user(b.id, f"SELECT 1 FROM {table}") == []  # noqa: S608
        assert await as_user(lead.id, f"SELECT 1 FROM {table}") == []  # noqa: S608
        assert await as_user(None, f"SELECT 1 FROM {table}") == []  # noqa: S608
        assert await as_user(a.id, f"SELECT 1 FROM {table}"), table  # noqa: S608
        count = await as_migrator(pg.migrator_url, f"SELECT count(*) FROM {table}")  # noqa: S608
        assert count == [(0,)]

    # Removed from the workspace, the creator sees nothing either.
    async with transaction(context={"app.user_id": str(lead.id)}) as tx:
        await tx.execute(
            text("DELETE FROM space_members WHERE space_id = :s AND user_id = :u"),
            {"s": workspace, "u": a.id},
        )
    assert await as_user(a.id, "SELECT 1 FROM outputs") == []
    with pytest.raises(service.OutputNotFoundError):
        await service.get_output(await context_of(a), created.id)


# --- AC #4 and #5 over the API ------------------------------------------------------------


async def _link(client: httpx.AsyncClient, output_id: uuid.UUID) -> str:
    response = await client.post(f"{OUTPUTS}/{output_id}/download-links")
    assert response.status_code == 201
    assert response.headers["cache-control"] == "no-store"
    url: str = response.json()["url"]
    assert url.startswith("/api/v1/downloads/")
    return url


async def test_download_csv_watermarked_and_audited(auth_env: AuthEnv) -> None:
    user = await make_user("anika")
    lead = await make_user("lead")
    workspace = await make_workspace("Project Falcon", lead, user)
    created = await create(user, workspace, PLANTED, title="Falcon margins")
    async with signed_in("anika") as client:
        url = await _link(client, created.id)
        response = await client.get(url)
        again = await client.get(url)  # valid for 24 h, not single-use
    assert response.status_code == 200
    assert again.status_code == 200
    assert response.headers["content-type"] == "text/csv; charset=utf-8"
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["content-disposition"] == (
        "attachment; filename=\"Falcon margins.csv\"; filename*=UTF-8''Falcon%20margins.csv"
    )
    body = response.content
    assert body.startswith(PLANTED)
    last = body.rstrip(b"\r\n").rsplit(b"\r\n", 1)[-1].decode()
    assert last.startswith("# CONFIDENTIAL - downloaded by anika on ")
    assert last.endswith(" IST - Project Falcon")

    rows = await audit_rows("outputs.downloaded")
    assert len(rows) == 2
    assert rows[0].payload["kind"] == "csv"
    assert rows[0].payload["size_bytes"] == len(PLANTED)
    assert rows[0].payload["file_digest"] == KeyedHash.of(hashlib.sha256(PLANTED).digest())
    assert rows[0].actor_user_id == user.id
    for row in await audit_rows():
        assert_no_content(row.payload)
        assert "Falcon" not in str(row.payload)


@pytest.mark.parametrize("kind", [OutputKind.XLSX, OutputKind.DOCX])
async def test_download_xlsx_and_docx(auth_env: AuthEnv, kind: OutputKind) -> None:
    user = await make_user("bo")
    ctx = await context_of(user)
    data = (
        xlsx_bytes([["Year", "EBITDA"], ["FY24", 12.5]])
        if kind is OutputKind.XLSX
        else docx_bytes("Summary")
    )
    created = await create(user, private_space(ctx), data, kind=kind, title="Report")
    async with signed_in("bo") as client:
        response = await client.get(await _link(client, created.id))
    assert response.status_code == 200
    if kind is OutputKind.XLSX:
        workbook = load_workbook(io.BytesIO(response.content))
        assert workbook["Notes"]["A1"].value.endswith("- Private space")
        assert workbook["Sheet"]["B2"].value == 12.5
    else:
        document = Document(io.BytesIO(response.content))
        assert any("Private space" in p.text for p in document.sections[0].footer.paragraphs)


async def test_token_is_for_its_user_only(auth_env: AuthEnv) -> None:
    a = await make_user("alice")
    await make_user("mallory")
    created = await create(a, private_space(await context_of(a)), PLANTED)
    async with signed_in("alice") as client:
        url = await _link(client, created.id)
    async with signed_in("mallory") as client:
        response = await client.get(url)
    assert response.status_code == 404
    assert "Planted" not in response.text
    assert await audit_rows("outputs.downloaded") == []


async def test_expired_token_is_404(auth_env: AuthEnv) -> None:
    a = await make_user("alice")
    ctx = await context_of(a)
    created = await create(a, private_space(ctx), PLANTED)
    issued = await service.issue_download_link(ctx, created.id, now=utcnow() - timedelta(hours=25))
    async with signed_in("alice") as client:
        assert (await client.get(issued.url)).status_code == 404
        assert (await client.get("/api/v1/downloads/unknown-token")).status_code == 404


async def test_removed_member_loses_download(auth_env: AuthEnv) -> None:
    a = await make_user("alice")
    lead = await make_user("lead")
    workspace = await make_workspace("Falcon", lead, a)
    created = await create(a, workspace, PLANTED)
    async with signed_in("alice") as client:
        url = await _link(client, created.id)
        async with transaction(context={"app.user_id": str(lead.id)}) as tx:
            await tx.execute(
                text("DELETE FROM space_members WHERE space_id = :s AND user_id = :u"),
                {"s": workspace, "u": a.id},
            )
        assert (await client.get(url)).status_code == 404
        assert (await client.get(f"{OUTPUTS}/{created.id}")).status_code == 404


async def test_download_streams_large_csv() -> None:
    user = await make_user("eli")
    ctx = await context_of(user)
    data = b"x" * (5 * CHUNK_SIZE + 17)
    created = await create(user, private_space(ctx), data)
    issued = await service.issue_download_link(ctx, created.id)
    opened = await service.open_download(ctx, issued.token, username="eli")
    sizes = [len(c) async for c in opened.chunks]
    assert len(sizes) >= 6
    assert max(sizes) <= CHUNK_SIZE  # never more than one decrypted chunk at a time
    assert sum(sizes) > len(data)


async def test_damaged_file_aborts_the_stream(files_dir: Path) -> None:
    user = await make_user("fay")
    ctx = await context_of(user)
    space_id = private_space(ctx)
    created = await create(user, space_id, b"y" * (2 * CHUNK_SIZE))
    path = files_dir / str(space_id) / str(created.id)
    path.write_bytes(path.read_bytes()[:-10])  # truncated
    issued = await service.issue_download_link(ctx, created.id)
    opened = await service.open_download(ctx, issued.token, username="fay")
    with pytest.raises(IntegrityError):
        async for _ in opened.chunks:
            pass


async def test_list_detail_versions(auth_env: AuthEnv) -> None:
    user = await make_user("gus")
    ctx = await context_of(user)
    space_id = private_space(ctx)
    first = await create(user, space_id, PLANTED, title="Model v1")
    second = await create(user, space_id, PLANTED, title="Model v2", version_of_id=first.id)
    async with signed_in("gus") as client:
        listed = (await client.get(OUTPUTS)).json()["items"]
        filtered = (await client.get(OUTPUTS, params={"space_id": str(space_id)})).json()["items"]
        other = (await client.get(OUTPUTS, params={"space_id": str(uuid.uuid4())})).json()["items"]
        detail = (await client.get(f"{OUTPUTS}/{first.id}")).json()
    assert [o["title"] for o in listed] == ["Model v2", "Model v1"]
    assert filtered == listed
    assert other == []
    assert [v["id"] for v in detail["versions"]] == [str(second.id)]
    assert listed[0]["version_of_id"] == str(first.id)


async def test_delete_shreds_key_removes_file_and_row(auth_env: AuthEnv, files_dir: Path) -> None:
    user = await make_user("hal")
    ctx = await context_of(user)
    space_id = private_space(ctx)
    created = await create(user, space_id, PLANTED)
    path = files_dir / str(space_id) / str(created.id)
    stored = path.read_bytes()
    async with signed_in("hal") as client:
        url = await _link(client, created.id)
        assert (await client.delete(f"{OUTPUTS}/{created.id}")).status_code == 204
        assert (await client.get(f"{OUTPUTS}/{created.id}")).status_code == 404
        assert (await client.get(url)).status_code == 404
    assert not path.exists()
    ref = service.KeyRef(space_id, service.OBJECT_TYPE, created.id)
    with pytest.raises(KeyDestroyed):
        async for _ in decrypt_stream(ref, chunks(stored), name="content"):
            pass
    assert await as_user(user.id, "SELECT 1 FROM download_tokens") == []
    (row,) = await audit_rows("outputs.deleted")
    assert row.payload == {"output_id": str(created.id), "space_id": str(space_id)}


@pytest.mark.parametrize("role", [Role.ADMIN, Role.COMPLIANCE])
async def test_staff_have_no_outputs(auth_env: AuthEnv, role: Role) -> None:
    await make_user("staff", role=role)
    async with signed_in("staff") as client:
        assert (await client.get(OUTPUTS)).status_code == 403


def test_download_filename_is_safe() -> None:
    assert service.download_filename('../a/b:c*"d', OutputKind.CSV) == "a_b_c_d.csv"
    assert service.download_filename("...", OutputKind.XLSX) == "output.xlsx"
    assert service.download_filename("Fälke 鷹", OutputKind.DOCX) == "Fälke 鷹.docx"
