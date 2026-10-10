"""AC #3: watermarks on XLSX, DOCX and CSV, values-only."""

import io
from datetime import UTC, datetime

from docx import Document
from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]

from app.outputs import watermark
from app.outputs.watermark import Watermark
from tests.outputs.conftest import chunks, docx_bytes, xlsx_bytes

AT = datetime(2026, 10, 10, 6, 0, tzinfo=UTC)  # 11:30 IST
MARK = Watermark(username="anika", at=AT, code_name="Project Falcon")
EXPECTED = "CONFIDENTIAL - downloaded by anika on 2026-10-10 11:30 IST - Project Falcon"


def test_text_uses_ist_and_private_label() -> None:
    assert MARK.text == EXPECTED
    assert Watermark("bo", AT, None).text.endswith("- Private space")


def test_xlsx_notes_sheet_and_every_footer() -> None:
    # Stored outputs are values-only: formula-like text was prefixed when generated.
    risky = watermark.values_only("=cmd|' /C calc'!A0")
    data = xlsx_bytes([["Year", "Revenue"], ["FY24", 1250.0], [risky, 3]], sheets=2)
    out = load_workbook(io.BytesIO(watermark.watermark_xlsx(data, MARK)))
    assert out.sheetnames == ["Sheet", "Sheet2", "Notes"]
    assert out["Notes"]["A1"].value == EXPECTED
    for sheet in out.worksheets:
        assert sheet.oddFooter.center.text == EXPECTED
        assert sheet.evenFooter.center.text == EXPECTED
    # Values are kept as they were.
    assert out["Sheet"]["B2"].value == 1250.0
    assert out["Sheet"]["A3"].value == risky
    # No cell is a formula.
    assert all(
        cell.data_type != "f"
        for sheet in out.worksheets
        for row in sheet.iter_rows()
        for cell in row
    )


def test_xlsx_existing_notes_sheet_gets_a_new_line() -> None:
    workbook = Workbook()
    workbook.active.title = "Data"
    notes = workbook.create_sheet("Notes")
    notes["A1"] = "Built from FY24 report"
    buf = io.BytesIO()
    workbook.save(buf)
    out = load_workbook(io.BytesIO(watermark.watermark_xlsx(buf.getvalue(), MARK)))
    assert [c.value for c in out["Notes"]["A"]] == ["Built from FY24 report", EXPECTED]


def test_xlsx_formula_like_code_name_stays_text() -> None:
    mark = Watermark("anika", AT, '=HYPERLINK("http://x")&Co')
    out = load_workbook(io.BytesIO(watermark.watermark_xlsx(xlsx_bytes([["a"]]), mark)))
    cell = out["Notes"]["A1"]
    assert cell.data_type == "s"
    assert "=HYPERLINK" in cell.value
    # "&" is the footer control character: escaped as "&&".
    assert out["Sheet"].oddFooter.center.text.endswith("&&Co")


def test_values_only_prefixes_formula_starts() -> None:
    assert [watermark.values_only(v) for v in ["=1+1", "+1", "-1", "@x", "ok"]] == [
        "'=1+1",
        "'+1",
        "'-1",
        "'@x",
        "ok",
    ]


def test_docx_footer() -> None:
    out = Document(io.BytesIO(watermark.watermark_docx(docx_bytes("Body text"), MARK)))
    footer_text = [p.text for p in out.sections[0].footer.paragraphs]
    assert EXPECTED in footer_text
    assert out.paragraphs[0].text == "Body text"


def test_unicode_code_name_in_every_format() -> None:
    mark = Watermark("anika", AT, "Projekt Fälke — 鷹")
    xlsx = load_workbook(io.BytesIO(watermark.watermark_xlsx(xlsx_bytes([["a"]]), mark)))
    assert xlsx["Notes"]["A1"].value.endswith("Projekt Fälke — 鷹")
    docx = Document(io.BytesIO(watermark.watermark_docx(docx_bytes("x"), mark)))
    assert any("鷹" in p.text for p in docx.sections[0].footer.paragraphs)


async def _csv(data: bytes, mark: Watermark = MARK) -> bytes:
    return b"".join([c async for c in watermark.watermark_csv(chunks(data, 3), mark)])


async def test_csv_last_line() -> None:
    out = await _csv(b"a,b\r\n1,2\r\n")
    assert out == b"a,b\r\n1,2\r\n# " + EXPECTED.encode() + b"\r\n"


async def test_csv_without_trailing_newline_and_empty() -> None:
    assert (await _csv(b"a,b")).split(b"\r\n")[:2] == [b"a,b", b"# " + EXPECTED.encode()]
    assert await _csv(b"") == b"# " + EXPECTED.encode() + b"\r\n"
