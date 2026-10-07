"""Story 1.7 acceptance tests for the synthetic corpus."""

import shutil
from collections import Counter
from pathlib import Path

import pypdfium2 as pdfium  # type: ignore[import-untyped]
import pytest
from docx import Document
from openpyxl import load_workbook  # type: ignore[import-untyped]
from pptx import Presentation

import check_corpus
from generate.build import DEFAULT_MANIFEST, dump, generate, load
from generate.ids import gstin_check_char, isin_check_digit, luhn_valid, verhoeff_valid
from generate.manifest_schema import Manifest
from generate.verify import verify

CORPUS_DIR = Path(__file__).resolve().parents[1]


# --- AC #1: byte-stable, hashes recorded -------------------------------------------------


def test_two_builds_are_byte_identical(build_dir: Path, manifest: Manifest, tmp_path: Path) -> None:
    second = generate(tmp_path / "build")
    assert dump(second) == dump(manifest)
    for doc in manifest.documents:
        assert (build_dir / doc.file).read_bytes() == (tmp_path / "build" / doc.file).read_bytes()


def test_committed_manifest_matches_the_build(manifest: Manifest) -> None:
    assert DEFAULT_MANIFEST.read_text(encoding="utf-8") == dump(manifest), (
        "manifest.yaml is stale: run `make corpus` and commit it"
    )


def test_manifest_round_trips(manifest: Manifest) -> None:
    assert load(DEFAULT_MANIFEST) == manifest


# --- AC #2-#4: everything is where the manifest says --------------------------------------


def test_self_check_passes(build_dir: Path, manifest: Manifest) -> None:
    assert verify(build_dir, manifest) == []


def test_document_mix(build_dir: Path, manifest: Manifest) -> None:
    files = sorted(p.name for p in build_dir.iterdir())
    assert Counter(Path(f).suffix for f in files) == {
        ".pdf": 5,
        ".docx": 3,
        ".pptx": 1,
        ".xlsx": 3,
        ".csv": 1,
    }
    cim_pages = sorted(
        len(pdfium.PdfDocument(str(build_dir / f))) for f in files if f.startswith("cim_")
    )
    assert cim_pages == [60, 110, 150]
    for f in ("scan_kyc_form.pdf", "scan_board_resolution.pdf"):
        pdf = pdfium.PdfDocument(str(build_dir / f))
        assert all(not pdf[i].get_textpage().get_text_range().strip() for i in range(len(pdf)))

    deck = Presentation(str(build_dir / "management_presentation_sahyadri.pptx"))
    assert len(deck.slides) == 25
    assert all(s.notes_slide.notes_text_frame.text.strip() for s in deck.slides)

    spa = Document(str(build_dir / "spa_vardhan_chemicals.docx"))
    assert len(spa.element.body.xpath('.//w:br[@w:type="page"]')) + 1 == 40

    invoices = load_workbook(build_dir / "invoices_nilgiri_logistics.xlsx", read_only=True)
    assert invoices["Transactions"].max_row - 1 >= 50_000
    model = load_workbook(build_dir / "model_vardhan_chemicals.xlsx")
    assert model["Internal_Notes"].sheet_state == "hidden"
    for name in (
        "model_vardhan_chemicals.xlsx",
        "budget_sahyadri_healthcare.xlsx",
        "invoices_nilgiri_logistics.xlsx",
    ):
        assert {"P&L", "Segments", "KPIs"} <= set(
            load_workbook(build_dir / name, read_only=True).sheetnames
        )


def test_cached_and_uncached_formulas(build_dir: Path, manifest: Manifest) -> None:
    formulas = load_workbook(build_dir / "model_vardhan_chemicals.xlsx")
    values = load_workbook(build_dir / "model_vardhan_chemicals.xlsx", data_only=True)
    pl, cached = formulas["P&L"], values["P&L"]
    formula_cells = [
        c
        for row in pl.iter_rows()
        for c in row
        if isinstance(c.value, str) and c.value.startswith("=")
    ]
    assert len(formula_cells) >= 50
    assert all(isinstance(cached[c.coordinate].value, float) for c in formula_cells)
    uncached = [u for d in manifest.documents for u in d.uncached_formulas]
    assert uncached == ["KPIs!K6"]
    assert formulas["KPIs"]["K6"].value.startswith("=")
    assert values["KPIs"]["K6"].value is None


def test_planted_item_counts(manifest: Manifest) -> None:
    kinds = Counter(i.kind for i in manifest.items)
    assert kinds["identifier"] >= 60
    assert kinds["benign"] >= 40
    assert kinds["injection"] >= 20
    assert kinds["canary"] == len(manifest.documents)
    types = {i.type for i in manifest.items if i.kind == "identifier"}
    assert types == {
        "pan",
        "aadhaar",
        "aadhaar_masked",
        "bank_account",
        "demat_nsdl",
        "demat_cdsl",
        "card",
        "upi",
        "passport",
        "voter_id",
    }
    assert {i.type for i in manifest.items if i.kind == "injection"} == {
        "white_text",
        "tiny_text",
        "docx_comment",
        "speaker_note",
        "cell_comment",
        "metadata",
        "hidden_sheet",
    }
    assert {i.expected for i in manifest.items if i.kind == "identifier"} == {"mask"}
    assert {i.expected for i in manifest.items if i.kind == "benign"} == {"keep"}


def test_edge_cases(manifest: Manifest, build_dir: Path) -> None:
    gstins = [i for i in manifest.items if i.type == "gstin"]
    assert gstins
    assert all(i.kind == "benign" for i in gstins)
    assert all(i.value[2:12][3] == "C" for i in gstins)  # embedded company PAN
    splits = [i for i in manifest.items if i.variant == "split-line"]
    assert {i.doc.rsplit(".", 1)[1] for i in splits} == {"pdf", "docx"}
    assert all(i.kind == "identifier" and i.match == "normalized" for i in splits)
    rupee_docs = [
        (build_dir / "vendor_ledger_nilgiri.csv").read_text(encoding="utf-8"),
        "\n".join(
            c.text
            for t in Document(str(build_dir / "term_sheet_v3.docx")).tables
            for r in t.rows
            for c in r.cells
        ),
    ]
    assert all("₹" in text for text in rupee_docs)
    ws = load_workbook(build_dir / "invoices_nilgiri_logistics.xlsx", read_only=True)[
        "Transactions"
    ]
    assert "₹" in next(ws.iter_rows(max_row=1, values_only=True))[4]


def test_checksums_are_valid(manifest: Manifest) -> None:
    for item in manifest.items:
        digits = "".join(ch for ch in item.value if ch.isdigit())
        if item.type == "aadhaar":
            assert verhoeff_valid(digits), item.id
        elif item.type == "card":
            assert luhn_valid(digits), item.id
        elif item.type == "gstin":
            assert gstin_check_char(item.value[:14]) == item.value[14], item.id
        elif item.type == "isin":
            assert isin_check_digit(item.value[:11]) == item.value[11], item.id


def test_reference_facts(manifest: Manifest) -> None:
    for doc in manifest.documents:
        if doc.file.startswith("cim_"):
            assert len(doc.figures) >= 15
            assert {"company", "stake_offered", "advisor"} <= set(doc.key_terms)
    diffs = manifest.term_sheet_diff
    assert len(diffs) >= 12
    assert Counter(d.change for d in diffs) == {"changed": 10, "removed": 2, "added": 3}
    assert {t.doc for t in manifest.workbook_totals} == {
        "model_vardhan_chemicals.xlsx",
        "budget_sahyadri_healthcare.xlsx",
        "invoices_nilgiri_logistics.xlsx",
    }


# --- negative cases: the self-check really catches problems --------------------------------


@pytest.fixture
def copy(build_dir: Path, tmp_path: Path) -> Path:
    target = tmp_path / "build"
    shutil.copytree(build_dir, target)
    return target


def test_stray_file_fails_check(copy: Path, manifest: Manifest, tmp_path: Path) -> None:
    (copy / "real_deal_notes.pdf").write_bytes(b"%PDF-1.4 not allowed")
    manifest_path = tmp_path / "manifest.yaml"
    manifest_path.write_text(dump(manifest), encoding="utf-8")
    assert check_corpus.main(["--build", str(copy), "--manifest", str(manifest_path)]) == 1
    assert any("real_deal_notes.pdf" in p for p in verify(copy, manifest))


def test_missing_and_modified_files_fail(copy: Path, manifest: Manifest) -> None:
    (copy / "vendor_ledger_nilgiri.csv").unlink()
    assert any("missing from build" in p for p in verify(copy, manifest))
    shutil.copy(copy / "term_sheet_v3.docx", copy / "vendor_ledger_nilgiri.csv")
    assert any("SHA-256" in p for p in verify(copy, manifest))


def test_wrong_locator_or_total_fails(build_dir: Path, manifest: Manifest) -> None:
    broken = manifest.model_copy(deep=True)
    item = next(i for i in broken.items if i.type == "voter_id" and i.doc.startswith("cim_"))
    item.locator = "p.2"
    broken.workbook_totals[0].value += 1
    broken.term_sheet_diff[0].v4 = "A clause that is in neither version."
    problems = verify(build_dir, broken)
    assert any(item.id in p for p in problems)
    assert any("recomputes" in p for p in problems)
    assert any("TS-D01" in p for p in problems)


def test_ocr_items_are_not_in_a_text_layer(manifest: Manifest) -> None:
    ocr = [i for i in manifest.items if i.match == "ocr"]
    assert {i.doc for i in ocr} == {"scan_kyc_form.pdf", "scan_board_resolution.pdf"}


# --- AC #5: fabricated only -----------------------------------------------------------------


def test_readme_forbids_real_documents() -> None:
    readme = (CORPUS_DIR / "README.md").read_text(encoding="utf-8")
    assert "Real deal documents must never be added" in readme


def test_no_real_looking_domains(manifest: Manifest) -> None:
    emails = [i.value for i in manifest.items if i.type == "business_email"]
    assert emails
    assert all(e.endswith(".example") for e in emails)
