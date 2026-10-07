"""API mode: drive the deployed app over its documented HTTP API (architecture §6).

Flow per run: log in as the eval user, find the private space, upload each corpus
document once and wait until it is ready. Per case: create a conversation, set the
selected documents, send each turn and parse the SSE stream. Then poll background jobs,
apply operation previews for ``excel_op`` cases, and download outputs.

The app is built in later stories, so any endpoint answering 404 or 501 raises
:class:`FeatureUnavailable` and the case is reported ``skipped: feature-unavailable``.
Response shapes the contract leaves open (job results, conversation messages) are read
tolerantly; see :meth:`ApiTarget._job_outcome`.
"""

import io
import time
from collections.abc import Callable
from typing import Any, Final

import httpx

from evals.runner import corpus
from evals.runner.clients import FeatureUnavailable
from evals.runner.schema import Case
from evals.runner.sse import parse, reduce, table_from
from evals.runner.transcript import Source, Table, Transcript, TurnResult

API: Final = "/api/v1"
READY: Final = {"ready", "ready_with_warnings"}
FAILED_DOC: Final = {"failed", "rejected"}
JOB_DONE: Final = {"succeeded", "completed", "done"}
JOB_FAILED: Final = {"failed", "cancelled", "error"}
CSRF_COOKIE: Final = "__Host-csrf"


class ApiError(RuntimeError):
    pass


class ApiTarget:
    mode = "api"

    def __init__(
        self,
        base_url: str,
        username: str,
        password: str,
        *,
        verify: bool | str = True,
        timeout_s: float = 120.0,
        ready_timeout_s: float = 900.0,
        job_timeout_s: float = 1800.0,
        poll_s: float = 2.0,
        transport: httpx.BaseTransport | None = None,
        sleep: Callable[[float], None] = time.sleep,
    ) -> None:
        self.base_url = base_url.rstrip("/")
        self._credentials = {"username": username, "password": password}
        self._http = httpx.Client(
            base_url=self.base_url, verify=verify, timeout=timeout_s, transport=transport
        )
        self._ready_timeout, self._job_timeout, self._poll = ready_timeout_s, job_timeout_s, poll_s
        self._sleep = sleep
        self._logged_in = False
        self._space_id: str | None = None
        self._doc_ids: dict[str, str] = {}  # corpus file -> app document id
        self._files: dict[str, str] = {}  # app document id -> corpus file

    def describe(self) -> dict[str, str]:
        return {"mode": "api", "target": self.base_url}

    def close(self) -> None:
        self._http.close()

    # --- HTTP -------------------------------------------------------------------------

    def _request(self, method: str, path: str, **kwargs: Any) -> httpx.Response:
        headers = dict(kwargs.pop("headers", {}))
        csrf = self._http.cookies.get(CSRF_COOKIE)
        if csrf and method not in {"GET", "HEAD"}:
            headers["X-CSRF-Token"] = csrf
        response = self._http.request(method, path, headers=headers, **kwargs)
        self._check(response, f"{method} {path}")
        return response

    @staticmethod
    def _check(response: httpx.Response, what: str) -> None:
        if response.status_code in {404, 405, 501}:
            raise FeatureUnavailable(f"{what} -> {response.status_code}")
        if response.status_code >= 400:  # noqa: PLR2004
            code = ""
            if "json" in response.headers.get("content-type", ""):
                code = str(response.json().get("code", ""))
            raise ApiError(f"{what} -> {response.status_code} {code}".strip())

    def _json(self, method: str, path: str, **kwargs: Any) -> dict[str, Any]:
        data = self._request(method, path, **kwargs).json()
        return data if isinstance(data, dict) else {"items": data}

    # --- session, space, documents ------------------------------------------------------

    def _ensure_session(self) -> None:
        if self._logged_in:
            return
        self._request("POST", f"{API}/auth/login", json=self._credentials)
        spaces = self._json("GET", f"{API}/spaces").get("items", [])
        private = [s for s in spaces if s.get("kind") == "private"] or spaces
        if not private:
            raise ApiError("the eval user has no space")
        self._space_id = str(private[0]["id"])
        self._logged_in = True

    def _upload(self, file: str) -> str:
        if file in self._doc_ids:
            return self._doc_ids[file]
        path = corpus.document_path(file)
        with path.open("rb") as fh:
            doc = self._json(
                "POST", f"{API}/spaces/{self._space_id}/documents", files={"file": (file, fh)}
            )
        doc_id = str(doc["id"])
        deadline = time.monotonic() + self._ready_timeout
        status = doc.get("status")
        while status not in READY:
            if status in FAILED_DOC:
                raise ApiError(f"{file}: ingestion {status} ({doc.get('status_code')})")
            if time.monotonic() > deadline:
                raise ApiError(f"{file}: not ready after {self._ready_timeout:.0f}s")
            self._sleep(self._poll)
            doc = self._json("GET", f"{API}/documents/{doc_id}")
            status = doc.get("status")
        self._doc_ids[file], self._files[doc_id] = doc_id, file
        return doc_id

    # --- jobs and outputs ------------------------------------------------------------------

    def _wait_job(self, job_id: str) -> dict[str, Any]:
        deadline = time.monotonic() + self._job_timeout
        while True:
            job = self._json("GET", f"{API}/jobs/{job_id}")
            status = str(job.get("status", ""))
            if status in JOB_DONE:
                return job
            if status in JOB_FAILED:
                raise ApiError(f"job {job_id} {status} ({job.get('error_code')})")
            if time.monotonic() > deadline:
                raise ApiError(f"job {job_id} timed out after {self._job_timeout:.0f}s")
            self._sleep(self._poll)

    def _download_output(self, output_id: str) -> Table:
        link = self._json("POST", f"{API}/outputs/{output_id}/download-links")
        url = link.get("url") or f"{API}/downloads/{link['token']}"
        content = self._request("GET", str(url)).content
        from openpyxl import load_workbook  # type: ignore[import-untyped]  # noqa: PLC0415

        ws = load_workbook(io.BytesIO(content), data_only=True, read_only=True).worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        header = [str(c) for c in rows[0]] if rows else []
        return Table(columns=header, rows=rows[1:], title=output_id, origin="output")

    def _job_outcome(self, conversation_id: str, job: dict[str, Any], turn: TurnResult) -> None:
        """Fold a finished job into the turn: an output file and/or the final message."""
        result = job.get("result") or {}
        output_id = result.get("output_id") or job.get("output_id")
        if output_id:
            turn.tables.append(self._download_output(str(output_id)))
        conversation = self._json("GET", f"{API}/conversations/{conversation_id}")
        messages = [m for m in conversation.get("messages", []) if m.get("role") == "assistant"]
        if messages:
            last = messages[-1]
            turn.text = str(last.get("content") or last.get("text") or turn.text)
            for item in last.get("sources", []):
                doc = str(item.get("document_id", ""))
                turn.sources.append(Source(self._files.get(doc, doc), str(item.get("locator", ""))))

    # --- a case ------------------------------------------------------------------------------

    def run_case(self, case: Case) -> Transcript:
        self._ensure_session()
        doc_ids = [self._upload(f) for f in case.documents]
        conversation = self._json("POST", f"{API}/conversations", json={"space_id": self._space_id})
        cid = str(conversation["id"])
        if doc_ids:
            self._request(
                "PUT", f"{API}/conversations/{cid}/documents", json={"document_ids": doc_ids}
            )
        transcript = Transcript(prompt_documents=list(case.documents))
        for text in case.turns:
            turn = self._send(cid, text)
            for job in turn.jobs:
                self._job_outcome(cid, self._wait_job(str(job["job_id"])), turn)
            if case.category == "excel_op":
                for plan in turn.plans:
                    applied = self._json(
                        "POST", f"{API}/operation-previews/{plan['preview_id']}/applications"
                    )
                    self._job_outcome(cid, self._wait_job(str(applied["job_id"])), turn)
            transcript.turns.append(turn)
        return transcript

    def _send(self, conversation_id: str, text: str) -> TurnResult:
        started = time.monotonic()
        path = f"{API}/conversations/{conversation_id}/messages"
        headers = {"Accept": "text/event-stream"}
        csrf = self._http.cookies.get(CSRF_COOKIE)
        if csrf:
            headers["X-CSRF-Token"] = csrf
        with self._http.stream("POST", path, json={"content": text}, headers=headers) as response:
            if response.status_code >= 400:  # noqa: PLR2004
                response.read()
                self._check(response, f"POST {path}")
            turn = reduce(parse(response.iter_lines()), lambda d: self._files.get(d, d))
        turn.latency_ms = turn.latency_ms or round((time.monotonic() - started) * 1000)
        for table in turn.tables:
            table.origin = "event"
        return turn


__all__ = ["ApiError", "ApiTarget", "table_from"]
