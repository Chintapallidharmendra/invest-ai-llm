"""Access to the synthetic corpus (Story 1.7): manifest, document text, masking.

Document text is split into *units* with the same locators the manifest uses
(``p.N``, ``para N``, ``slide N``, ``Sheet!A1:G40``, CSV ``A1:F40``), so model-mode
excerpts can be cited and scored like API citations.

Masking: model mode must never send raw planted identifiers (ADR-027). When the app's
recognizers (Story 4.1, ``app.guardrails.identifiers.mask``) are importable they are
used; otherwise every identifier the manifest says was planted is replaced with
``[TYPE]``. That is exact for this corpus.
"""

import csv
import importlib
import re
import sys
from collections.abc import Callable
from dataclasses import dataclass
from functools import cache
from pathlib import Path
from typing import Any, Final

from evals.runner.schema import REPO_ROOT

CORPUS_DIR: Final = REPO_ROOT / "evals" / "corpus"
BUILD_DIR: Final = CORPUS_DIR / "build"
SHEET_CHUNK_ROWS: Final = 40
MAX_SHEET_ROWS: Final = 120  # excerpts never inline more of a sheet than this

if str(CORPUS_DIR) not in sys.path:
    sys.path.insert(0, str(CORPUS_DIR))

from generate.manifest_schema import Manifest  # noqa: E402


class CorpusMissingError(RuntimeError):
    """``evals/corpus/build`` is absent: run ``make corpus`` first."""


@cache
def manifest() -> Manifest:
    from generate.build import load  # noqa: PLC0415 (heavy imports only when needed)

    return load(CORPUS_DIR / "manifest.yaml")


def document_path(file: str) -> Path:
    path = BUILD_DIR / file
    if not path.is_file():
        raise CorpusMissingError(f"{path} not found: run `make corpus` first")
    return path


@dataclass(frozen=True)
class Unit:
    locator: str
    text: str


def _pdf_units(path: Path) -> list[Unit]:
    import pypdfium2 as pdfium  # type: ignore[import-untyped]  # noqa: PLC0415

    doc = pdfium.PdfDocument(str(path))
    try:
        return [
            Unit(f"p.{i + 1}", str(doc[i].get_textpage().get_text_range())) for i in range(len(doc))
        ]
    finally:
        doc.close()


def _docx_units(path: Path) -> list[Unit]:
    from docx import Document  # noqa: PLC0415

    doc = Document(str(path))
    units = [
        Unit(f"para {i}", p.text) for i, p in enumerate(doc.paragraphs, start=1) if p.text.strip()
    ]
    for t, table in enumerate(doc.tables, start=1):
        rows = [" | ".join(c.text for c in row.cells) for row in table.rows]
        units.append(Unit(f"table {t}", "\n".join(rows)))
    return units


def _pptx_units(path: Path) -> list[Unit]:
    from pptx import Presentation  # noqa: PLC0415

    units = []
    for n, slide in enumerate(Presentation(str(path)).slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                texts.extend(" | ".join(c.text for c in row.cells) for row in shape.table.rows)
        units.append(Unit(f"slide {n}", "\n".join(texts)))
    return units


def _fmt_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:.10g}"
    return str(value)


def _rows_units(name: str, rows: list[list[Any]], prefix: str) -> list[Unit]:
    from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]  # noqa: PLC0415

    units = []
    width = max((len(r) for r in rows), default=1)
    for start in range(0, min(len(rows), MAX_SHEET_ROWS), SHEET_CHUNK_ROWS):
        chunk = rows[start : start + SHEET_CHUNK_ROWS]
        ref = f"A{start + 1}:{get_column_letter(width)}{start + len(chunk)}"
        text = "\n".join(",".join(_fmt_cell(v) for v in row) for row in chunk)
        if start + SHEET_CHUNK_ROWS >= MAX_SHEET_ROWS and len(rows) > MAX_SHEET_ROWS:
            text += f"\n... ({len(rows) - MAX_SHEET_ROWS} more rows not shown)"
        units.append(Unit(f"{prefix}{ref}", f"Sheet {name}:\n{text}"))
    return units


def _xlsx_units(path: Path) -> list[Unit]:
    from openpyxl import load_workbook  # type: ignore[import-untyped]  # noqa: PLC0415

    wb = load_workbook(path, data_only=True, read_only=True)  # cached values only
    units = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue  # hidden sheets are flagged content, not context
        rows = [
            list(r)
            for _, r in zip(range(MAX_SHEET_ROWS), ws.iter_rows(values_only=True), strict=False)
        ]
        units.extend(_rows_units(ws.title, rows, f"{ws.title}!"))
    wb.close()
    return units


def _csv_units(path: Path) -> list[Unit]:
    with path.open(encoding="utf-8", newline="") as fh:
        rows: list[list[Any]] = list(csv.reader(fh))
    return _rows_units(path.stem, rows, "")


_EXTRACTORS: Final[dict[str, Callable[[Path], list[Unit]]]] = {
    ".pdf": _pdf_units,
    ".docx": _docx_units,
    ".pptx": _pptx_units,
    ".xlsx": _xlsx_units,
    ".csv": _csv_units,
}


@cache
def units(file: str) -> tuple[Unit, ...]:
    path = document_path(file)
    return tuple(_EXTRACTORS[path.suffix](path))


@cache
def full_text(file: str) -> str:
    return "\n".join(u.text for u in units(file))


# --- masking --------------------------------------------------------------------------------


@cache
def _app_masker() -> Callable[[str], str] | None:
    backend = REPO_ROOT / "backend"
    if str(backend) not in sys.path:
        sys.path.insert(0, str(backend))
    try:
        identifiers = importlib.import_module("app.guardrails.identifiers")
    except ImportError:
        return None
    return lambda text: str(identifiers.mask(text).text)


@cache
def _manifest_patterns() -> tuple[tuple[re.Pattern[str], str], ...]:
    patterns = []
    for item in manifest().items:
        if item.kind != "identifier":
            continue
        chars = [re.escape(ch) for ch in item.value if not ch.isspace()]
        # Tolerate any whitespace (including a line break) between characters.
        patterns.append((re.compile(r"\s*".join(chars), re.IGNORECASE), f"[{item.type.upper()}]"))
    # Longest first, so a full number is masked before a substring of it.
    return tuple(sorted(patterns, key=lambda p: -len(p[0].pattern)))


def masking_method() -> str:
    return "app.guardrails.identifiers" if _app_masker() else "manifest planted values"


def mask(text: str) -> str:
    masker = _app_masker()
    if masker is not None:
        return masker(text)
    for pattern, placeholder in _manifest_patterns():
        text = pattern.sub(placeholder, text)
    return text
