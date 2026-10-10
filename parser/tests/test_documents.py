"""AC #2: PDF/DOCX/PPTX text, tables and locators against the corpus' planted facts, and
the page/slide limits of AC #4.

PDF tests need Docling's layout/TableFormer models at ``DOCLING_ARTIFACTS_PATH``;
without it they skip.
"""

import io
import zipfile
from typing import Any

import pytest
from docx import Document
from pptx import Presentation

from app.docling_parse import OLE_MAGIC, artifacts_path, parse_document
from app.schema import ErrorCode, FileKind, Limits, ParseError, ParseResult
from tests.conftest import corpus_file, document, pdf_bytes

LIMITS = Limits()

needs_models = pytest.mark.skipif(
    artifacts_path() is None or not artifacts_path().is_dir(),  # type: ignore[union-attr]
    reason="DOCLING_ARTIFACTS_PATH not set (Docling layout/TableFormer models)",
)

_cache: dict[str, ParseResult] = {}


def parsed(name: str, kind: FileKind) -> ParseResult:
    if name not in _cache:
        _cache[name] = parse_document(kind, corpus_file(name), LIMITS)
    return _cache[name]


def texts_at(result: ParseResult, locator: str) -> str:
    """All text (items and table cells) at ``locator``, whitespace-normalised."""
    parts = [t.text for t in result.text if t.locator == locator]
    parts += [c.text for t in result.tables if t.locator == locator for c in t.cells]
    return " ".join(" ".join(parts).split())


def found(item: dict[str, Any], result: ParseResult) -> bool:
    """The planted value is at its locator (``match: normalized``: ignoring whitespace)."""
    text = texts_at(result, item["locator"])
    if item.get("match") == "normalized":
        return "".join(item["value"].split()) in "".join(text.split())
    return bool(item["value"] in text)


def planted(manifest: dict[str, Any], name: str, prefixes: tuple[str, ...]) -> list[dict[str, Any]]:
    return [
        i
        for i in manifest["items"]
        if i["doc"] == name and i["locator"].startswith(prefixes) and i["type"] != "white_text"
    ]


# --- DOCX ------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    "name", ["term_sheet_v3.docx", "term_sheet_v4.docx", "spa_vardhan_chemicals.docx"]
)
def test_docx_planted_facts_at_their_paragraphs(manifest: dict[str, Any], name: str) -> None:
    result = parsed(name, FileKind.DOCX)
    items = planted(manifest, name, ("para ", "comment ", "meta:"))
    assert items
    for item in items:
        assert found(item, result), item["id"]
    expected = document(manifest, name)
    assert result.pages == expected["pages"]
    assert any(expected["canary"] in t.text for t in result.text)


def test_docx_kinds_and_tables(manifest: dict[str, Any]) -> None:
    result = parsed("term_sheet_v3.docx", FileKind.DOCX)
    kinds = {t.kind.value for t in result.text}
    assert {"heading", "text", "comment", "metadata"} <= kinds
    orders = sorted([t.order for t in result.text] + [t.order for t in result.tables])
    assert orders == list(range(len(orders)))  # one reading order across text and tables
    (table,) = result.tables
    assert table.locator == "table 1"
    assert table.n_cols == 2
    assert table.cells[0].text == "Term"
    assert table.cells[0].header


def test_docx_nested_tables_parse() -> None:
    doc = Document()
    outer = doc.add_table(rows=2, cols=2)
    inner = outer.cell(0, 0).add_table(rows=2, cols=2)
    inner.cell(1, 1).text = "deep value"
    outer.cell(1, 1).text = "outer value"
    out = io.BytesIO()
    doc.save(out)
    result = parse_document(FileKind.DOCX, out.getvalue(), LIMITS)
    cells = " ".join(c.text for t in result.tables for c in t.cells)
    assert "outer value" in cells


def test_docx_page_limit() -> None:
    doc = Document()
    for _ in range(5):
        doc.add_paragraph("page")
        doc.add_page_break()  # type: ignore[no-untyped-call]
    out = io.BytesIO()
    doc.save(out)
    with pytest.raises(ParseError) as excinfo:
        parse_document(FileKind.DOCX, out.getvalue(), Limits(max_pages=5))
    assert (excinfo.value.code, excinfo.value.detail) == (ErrorCode.LIMIT_EXCEEDED, "pages")


# --- PPTX ------------------------------------------------------------------------------------


def test_pptx_planted_facts_notes_and_figures(manifest: dict[str, Any]) -> None:
    name = "management_presentation_sahyadri.pptx"
    result = parsed(name, FileKind.PPTX)
    expected = document(manifest, name)
    assert result.slides == expected["slides"]
    for item in planted(manifest, name, ("slide ", "meta:")):
        assert found(item, result), item["id"]
    for figure in expected["figures"]:
        assert figure["text"] in texts_at(result, figure["locator"]), figure["label"]
    assert any(t.kind.value == "note" for t in result.text)
    assert all(t.locator.startswith(("slide ", "meta:")) for t in result.text)


def test_pptx_slide_limit() -> None:
    deck = Presentation()
    for _ in range(4):
        deck.slides.add_slide(deck.slide_layouts[5])
    out = io.BytesIO()
    deck.save(out)
    with pytest.raises(ParseError) as excinfo:
        parse_document(FileKind.PPTX, out.getvalue(), Limits(max_slides=3))
    assert (excinfo.value.code, excinfo.value.detail) == (ErrorCode.LIMIT_EXCEEDED, "slides")


# --- PDF -------------------------------------------------------------------------------------


@needs_models
def test_pdf_figures_at_their_pages(manifest: dict[str, Any]) -> None:
    name = "cim_project_alpha_vardhan_chemicals.pdf"
    result = parsed(name, FileKind.PDF)
    expected = document(manifest, name)
    assert result.pages == expected["pages"]
    assert result.needs_ocr == []
    for figure in expected["figures"]:
        assert figure["text"] in texts_at(result, figure["locator"]), figure["label"]
    for item in planted(manifest, name, ("p.", "meta:")):
        assert found(item, result), item["id"]


@needs_models
def test_pdf_tables_have_cells_and_locators() -> None:
    result = parsed("cim_project_alpha_vardhan_chemicals.pdf", FileKind.PDF)
    assert result.tables
    for table in result.tables:
        assert table.locator.startswith("p.")
        assert table.n_rows > 1
        assert table.n_cols > 1
        assert all(0 <= c.row < table.n_rows and 0 <= c.col < table.n_cols for c in table.cells)
    # The FY2022-FY2024 financials table on p.33 holds the planted totals.
    page_33 = texts_at(result, "p.33")
    assert "1,230.9" in page_33
    assert "1,406.0" in page_33


@needs_models
def test_image_only_pages_need_ocr(manifest: dict[str, Any]) -> None:
    for name in ("scan_kyc_form.pdf", "scan_board_resolution.pdf"):
        result = parsed(name, FileKind.PDF)
        pages = document(manifest, name)["pages"]
        assert result.needs_ocr == [f"p.{n}" for n in range(1, pages + 1)]


def test_pdf_page_limit_is_checked_before_conversion() -> None:
    with pytest.raises(ParseError) as excinfo:
        parse_document(FileKind.PDF, pdf_bytes(301), LIMITS)
    assert (excinfo.value.code, excinfo.value.detail) == (ErrorCode.LIMIT_EXCEEDED, "pages")


# --- Broken and hostile input ------------------------------------------------------------------


@pytest.mark.parametrize("kind", [FileKind.PDF, FileKind.DOCX, FileKind.PPTX])
@pytest.mark.parametrize("data", [b"", b"\x00" * 100, b"hello world", b"PK\x03\x04broken"])
def test_garbage_is_corrupt(kind: FileKind, data: bytes) -> None:
    with pytest.raises(ParseError) as excinfo:
        parse_document(kind, data, LIMITS)
    assert excinfo.value.code is ErrorCode.CORRUPT_FILE


@pytest.mark.parametrize("kind", [FileKind.DOCX, FileKind.PPTX])
def test_legacy_or_encrypted_office_is_unsupported(kind: FileKind) -> None:
    with pytest.raises(ParseError) as excinfo:
        parse_document(kind, OLE_MAGIC + b"\x00" * 512, LIMITS)
    assert excinfo.value.code is ErrorCode.UNSUPPORTED


def test_renamed_extension_fails_cleanly() -> None:
    xlsx_as_docx = io.BytesIO()
    with zipfile.ZipFile(xlsx_as_docx, "w") as package:
        package.writestr("xl/workbook.xml", "<workbook/>")
    with pytest.raises(ParseError) as excinfo:
        parse_document(FileKind.DOCX, xlsx_as_docx.getvalue(), LIMITS)
    assert excinfo.value.code is ErrorCode.CORRUPT_FILE
    with pytest.raises(ParseError):
        parse_document(FileKind.PPTX, corpus_file("term_sheet_v3.docx"), LIMITS)
