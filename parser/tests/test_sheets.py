"""AC #3 (XLSX, CSV) and the sheet/row limits of AC #4."""

import codecs
import io
import zipfile
from datetime import datetime
from typing import Any

import pytest
from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.comments import Comment  # type: ignore[import-untyped]
from openpyxl.utils.cell import (  # type: ignore[import-untyped]
    column_index_from_string,
    coordinate_from_string,
)

from app.csv_parse import parse_csv, sniff
from app.schema import ErrorCode, Limits, ParseError
from app.xlsx_parse import parse_xlsx
from tests.conftest import corpus_file, document

LIMITS = Limits()


def workbook_bytes(build: Any) -> bytes:
    workbook = Workbook()
    build(workbook)
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def cell(result: Any, sheet: str, ref: str) -> Any:
    letters, row = coordinate_from_string(ref)
    col = column_index_from_string(letters)
    (found,) = [s for s in result.sheets if s.name == sheet]
    for r in found.rows:
        if r.row == row:
            return r.values[col - 1] if col <= len(r.values) else None
    return None


# --- XLSX: corpus --------------------------------------------------------------------------


def test_corpus_model_workbook(manifest: dict[str, Any]) -> None:
    name = "model_vardhan_chemicals.xlsx"
    expected = document(manifest, name)
    result = parse_xlsx(corpus_file(name), LIMITS)
    assert [s.name for s in result.sheets] == expected["sheets"]
    assert [s.name for s in result.sheets if s.hidden] == expected["hidden_sheets"]
    # The one formula saved without a cached value (KPIs!K6) is counted, not evaluated.
    counts = {s.name: s.formulas_without_cached_value for s in result.sheets}
    assert counts["KPIs"] == 1
    assert sum(counts.values()) == len(expected["uncached_formulas"])
    assert cell(result, "KPIs", "K6") is None
    # Cached totals are returned as numbers (P&L!F5 is the FY2024 total revenue).
    (total,) = [
        i for i in manifest["workbook_totals"] if i["doc"] == name and i.get("cell") == "P&L!F5"
    ]
    assert cell(result, "P&L", "F5") == pytest.approx(total["value"])


def test_corpus_large_workbook_rows(manifest: dict[str, Any]) -> None:
    name = "invoices_nilgiri_logistics.xlsx"
    result = parse_xlsx(corpus_file(name), LIMITS)
    # The manifest records each sheet's last row number.
    last_rows = {s.name: s.rows[-1].row for s in result.sheets}
    assert last_rows == document(manifest, name)["rows"]
    count = next(
        i for i in manifest["workbook_totals"] if i["doc"] == name and i.get("method") == "count"
    )
    assert cell(result, "Summary", "B4") == count["value"]


def test_row_limit_on_corpus_workbook() -> None:
    data = corpus_file("invoices_nilgiri_logistics.xlsx")
    with pytest.raises(ParseError) as excinfo:
        parse_xlsx(data, Limits(max_rows_per_sheet=50_000))
    assert excinfo.value.code is ErrorCode.LIMIT_EXCEEDED


# --- XLSX: generated -----------------------------------------------------------------------


def test_values_dates_comments_and_no_evaluation() -> None:
    def build(workbook: Workbook) -> None:
        sheet = workbook.active
        sheet.title = "Data"
        sheet["A1"] = "Year"
        sheet["B1"] = 1250.5
        sheet["C1"] = 42
        sheet["D1"] = datetime(2024, 3, 31, 0, 0)
        sheet["E1"] = True
        sheet["F1"] = "=B1*2"  # no cached value: openpyxl never calculates
        sheet["A1"].comment = Comment("Checked by audit", "Asha")
        hidden = workbook.create_sheet("Hidden")
        hidden.sheet_state = "hidden"
        hidden["A3"] = "secret row"

    result = parse_xlsx(workbook_bytes(build), LIMITS)
    data, hidden = result.sheets
    assert data.rows[0].row == 1
    assert data.rows[0].values == ["Year", 1250.5, 42, "2024-03-31T00:00:00", True]
    assert data.formulas_without_cached_value == 1
    assert [(c.ref, c.text, c.author) for c in data.comments] == [
        ("A1", "Checked by audit", "Asha")
    ]
    assert hidden.hidden
    assert [(r.row, r.values) for r in hidden.rows] == [(3, ["secret row"])]


def test_external_links_are_not_followed() -> None:
    def build(workbook: Workbook) -> None:
        workbook.active["A1"] = "='[other.xlsx]Sheet1'!A1"

    data = workbook_bytes(build)
    result = parse_xlsx(data, LIMITS)
    (sheet,) = result.sheets
    assert sheet.formulas_without_cached_value == 1
    assert sheet.rows == []  # the linked value isn't fetched or computed


def test_sheet_limit() -> None:
    def build(workbook: Workbook) -> None:
        for i in range(5):
            workbook.create_sheet(f"S{i}")

    with pytest.raises(ParseError) as excinfo:
        parse_xlsx(workbook_bytes(build), Limits(max_sheets=5))
    assert (excinfo.value.code, excinfo.value.detail) == (ErrorCode.LIMIT_EXCEEDED, "sheets")


@pytest.mark.parametrize("data", [b"not a zip", b"PK\x03\x04garbage", b""])
def test_corrupt_xlsx(data: bytes) -> None:
    with pytest.raises(ParseError) as excinfo:
        parse_xlsx(data, LIMITS)
    assert excinfo.value.code is ErrorCode.CORRUPT_FILE


def test_zip_without_workbook_is_corrupt() -> None:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as package:
        package.writestr("hello.txt", "hi")
    with pytest.raises(ParseError) as excinfo:
        parse_xlsx(out.getvalue(), LIMITS)
    assert excinfo.value.code is ErrorCode.CORRUPT_FILE


# --- CSV -----------------------------------------------------------------------------------


def test_corpus_csv(manifest: dict[str, Any]) -> None:
    name = "vendor_ledger_nilgiri.csv"
    result = parse_csv(corpus_file(name), LIMITS)
    assert result.csv is not None
    assert result.csv.encoding == "utf-8"
    assert len(result.csv.rows) == document(manifest, name)["rows"]["csv"]  # data rows


@pytest.mark.parametrize(
    ("data", "encoding"),
    [
        ('Name,Amount\nFälke,"1,579.4"\n'.encode(), "utf-8"),
        (codecs.BOM_UTF8 + "Name,Amount\nFälke,1\n".encode(), "utf-8"),
        (codecs.BOM_UTF16_LE + "Name,Amount\nFälke,1\n".encode("utf-16-le"), "utf-16"),
        (codecs.BOM_UTF16_BE + "Name,Amount\nFälke,1\n".encode("utf-16-be"), "utf-16"),
        ("Name,Amount\nFälke,1\n".encode("cp1252"), "cp1252"),
    ],
)
def test_csv_encodings(data: bytes, encoding: str) -> None:
    result = parse_csv(data, LIMITS)
    assert result.csv is not None
    assert result.csv.encoding == encoding
    assert result.csv.columns == ["Name", "Amount"]
    assert result.csv.rows[0][0] == "Fälke"


def test_csv_values_stay_text() -> None:
    result = parse_csv(b'id,amount\n007,"1,579.4"\n,=SUM(A1)\n', LIMITS)
    assert result.csv is not None
    assert result.csv.rows == [["007", "1,579.4"], [None, "=SUM(A1)"]]


def test_csv_row_limit() -> None:
    data = b"a\n" + b"1\n" * 11
    with pytest.raises(ParseError) as excinfo:
        parse_csv(data, Limits(max_rows_per_sheet=10))
    assert excinfo.value.code is ErrorCode.LIMIT_EXCEEDED
    assert parse_csv(b"a\n" + b"1\n" * 10, Limits(max_rows_per_sheet=10)).csv is not None


def test_sniff_prefers_utf8() -> None:
    assert sniff("₹ 100".encode())[0] == "utf-8"
